import json
import os
import csv
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path
from tkinter import messagebox, StringVar, filedialog

import ttkbootstrap as ttk
from ttkbootstrap.constants import *


HISTORY_FILE = Path("config") / "history.json"


class HistoryPage:

    def __init__(self, parent, app):
        self.app = app

        self.frame = ttk.Frame(parent)
        self.frame.pack(fill=BOTH, expand=True)

        self.create_header()
        self.create_summary()
        self.create_table()
        self.create_buttons()

        self.load_history()

# ==================================================
# HEADER
# ==================================================

    def create_header(self):

        header = ttk.Frame(self.frame)
        header.pack(fill=X, pady=(0, 20))

        ttk.Label(
            header,
            text="📜 Automation History",
            font=("Segoe UI", 28, "bold"),
            bootstyle="primary"
        ).pack(side=LEFT)

        # Search variable
        self.search_var = StringVar()

        # Search box
        self.search_entry = ttk.Entry(
            header,
            textvariable=self.search_var,
            width=30
        )

        self.search_entry.pack(
            side=LEFT,
            padx=(40, 5)
        )

        self.search_entry.insert(
            0,
            "Search history..."
        )

        self.search_entry.bind(
            "<FocusIn>",
            self.clear_search_placeholder
        )

        self.search_entry.bind(
            "<KeyRelease>",
            self.filter_history
        )

        # Search button
        ttk.Button(
            header,
            text="🔍 Search",
            bootstyle="info",
            command=self.filter_history
        ).pack(
            side=LEFT,
            padx=5
        )

        # Refresh button
        ttk.Button(
            header,
            text="🔄 Refresh",
            bootstyle="secondary",
            command=self.load_history
        ).pack(
            side=RIGHT
        )
        ttk.Button(
            header,
            text="📥 Export CSV",
            command=self.export_history_csv,
            bootstyle="success"
        ).pack(
            side="left",
            padx=(10, 0)
        )
        
        ttk.Button(
            header,
            text="📊 Export Excel",
            command=self.export_history_excel,
            bootstyle="primary"
        ).pack(
            side="left",
            padx=5
        )
        
# ==================================================
# SEARCH
# ==================================================

    def clear_search_placeholder(self, event=None):

        if self.search_entry.get() == "Search history...":

            self.search_entry.delete(
                0,
                END
            )

    def filter_history(self, event=None):

        search_text = (
            self.search_var.get()
            .strip()
            .lower()
        )

        if search_text == "search history...":
            search_text = ""

        # Clear current table
        for item in self.tree.get_children():
            self.tree.delete(item)

        # History file doesn't exist
        if not HISTORY_FILE.exists():
            self.update_summary([])
            return

        try:

            with open(
                HISTORY_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                history = json.load(file)

        except Exception as e:

            messagebox.showerror(
                "History Error",
                f"Unable to load history:\n\n{e}"
            )

            return

        filtered_history = []

        for record in reversed(history):

            searchable_text = " ".join([
                str(record.get("date_time", "")),
                str(record.get("input_folder", "")),
                str(record.get("output_folder", "")),
                str(record.get("status", "")),
                str(record.get("excel_files", "")),
                str(record.get("total_rows", "")),
                str(record.get("duplicates_removed", "")),
                str(record.get("blank_rows_removed", "")),
                str(record.get("final_rows", ""))
            ]).lower()

            if search_text in searchable_text:

                filtered_history.append(record)

                self.insert_record(record)

        # Update summary based on filtered results
        self.update_summary(filtered_history)

# ==================================================
# SUMMARY
# ==================================================

    def create_summary(self):

        summary_frame = ttk.Frame(self.frame)

        summary_frame.pack(
            fill=X,
            pady=(0, 20)
        )

        self.total_runs_label = self.create_summary_card(
            summary_frame,
            "Total Runs",
            "0",
            "primary"
        )

        self.success_label = self.create_summary_card(
            summary_frame,
            "Successful",
            "0",
            "success"
        )

        self.files_label = self.create_summary_card(
            summary_frame,
            "Excel Files",
            "0",
            "info"
        )

        self.rows_label = self.create_summary_card(
            summary_frame,
            "Rows Processed",
            "0",
            "warning"
        )

# ==================================================
# ANALYTICS
# ==================================================

        analytics_frame = ttk.Frame(self.frame)

        analytics_frame.pack(
            fill=X,
            pady=(0, 20)
        )

        self.avg_time_label = self.create_summary_card(
            analytics_frame,
            "Avg Processing Time",
            "0 sec",
            "secondary"
        )

        self.fastest_run_label = self.create_summary_card(
            analytics_frame,
            "Fastest Run",
            "0 sec",
            "success"
        )

        self.duplicates_label = self.create_summary_card(
            analytics_frame,
            "Duplicates Removed",
            "0",
            "danger"
        )
    

    def create_summary_card(
        self,
        parent,
        title,
        value,
        style
    ):

        card = ttk.Frame(
            parent,
            padding=15,
            bootstyle=style
        )

        card.pack(
            side=LEFT,
            fill=X,
            expand=True,
            padx=5
        )

        ttk.Label(
            card,
            text=title,
            font=("Segoe UI", 11)
        ).pack()

        label = ttk.Label(
            card,
            text=value,
            font=("Segoe UI", 22, "bold")
        )

        label.pack(
            pady=(5, 0)
        )

        return label

    # ==================================================
    # TABLE
    # ==================================================

    def create_table(self):

        table_frame = ttk.Frame(self.frame)

        table_frame.pack(
            fill=BOTH,
            expand=True
        )

        columns = (
            "date",
            "input",
            "files",
            "rows",
            "duplicates",
            "blank",
            "final",
            "time",
            "status"
        )

        self.tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            bootstyle="primary"
        )

        headings = {
            "date": "Date & Time",
            "input": "Input Folder",
            "files": "Files",
            "rows": "Total Rows",
            "duplicates": "Duplicates",
            "blank": "Blank Rows",
            "final": "Final Rows",
            "time": "Processing Time",
            "status": "Status"
        }

        widths = {
            "date": 160,
            "input": 300,
            "files": 70,
            "rows": 100,
            "duplicates": 100,
            "blank": 100,
            "final": 100,
            "time": 120,
            "status": 100
        }

        for column in columns:

            self.tree.heading(
                column,
                text=headings[column]
            )

            self.tree.column(
                column,
                width=widths[column],
                anchor=CENTER
            )

        # Vertical scrollbar
        scrollbar_y = ttk.Scrollbar(
            table_frame,
            orient=VERTICAL,
            command=self.tree.yview
        )

        # Horizontal scrollbar
        scrollbar_x = ttk.Scrollbar(
            table_frame,
            orient=HORIZONTAL,
            command=self.tree.xview
        )

        self.tree.configure(
            yscrollcommand=scrollbar_y.set,
            xscrollcommand=scrollbar_x.set
        )

        self.tree.pack(
            side=LEFT,
            fill=BOTH,
            expand=True
        )

        scrollbar_y.pack(
            side=RIGHT,
            fill=Y
        )

        scrollbar_x.pack(
            side=BOTTOM,
            fill=X
        )
# Double-click a history record to view details
        self.tree.bind(
            "<Double-1>",
            self.show_run_details
        )

    # ==================================================
    # INSERT HISTORY RECORD
    # ==================================================

    def insert_record(self, record):

        self.tree.insert(
            "",
            END,
            values=(
                record.get(
                    "date_time",
                    ""
                ),

                record.get(
                    "input_folder",
                    ""
                ),

                record.get(
                    "excel_files",
                    0
                ),

                f"{record.get('total_rows', 0):,}",

                f"{record.get('duplicates_removed', 0):,}",

                f"{record.get('blank_rows_removed', 0):,}",

                f"{record.get('final_rows', 0):,}",

                f"{record.get('processing_time', 0):.2f}s",

                record.get(
                    "status",
                    "Unknown"
                )
            )
        )
        # ==================================================
    # VIEW RUN DETAILS
    # ==================================================

    def show_run_details(self, event=None):

        selected = self.tree.selection()

        if not selected:
            return

        item = self.tree.item(selected[0])

        values = item.get("values", [])

        if not values:
            return

        date_time = values[0]

        # Load complete history
        if not HISTORY_FILE.exists():
            return

        try:

            with open(
                HISTORY_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                history = json.load(file)

        except Exception as e:

            messagebox.showerror(
                "History Error",
                f"Unable to load history:\n\n{e}"
            )

            return

        # Find the selected record
        selected_record = None

        for record in history:

            if record.get("date_time", "") == date_time:

                selected_record = record
                break

        if selected_record is None:

            messagebox.showwarning(
                "Run Details",
                "Unable to find the selected automation run."
            )

            return

        # Create details window
        details_window = ttk.Toplevel(
            self.frame
        )

        details_window.title(
            "Automation Run Details"
        )

        details_window.geometry(
            "700x620"
        )

        details_window.resizable(
            False,
            False
        )

        # Header
        ttk.Label(
            details_window,
            text="📋 Automation Run Details",
            font=("Segoe UI", 22, "bold"),
            bootstyle="primary"
        ).pack(
            pady=(25, 20)
        )

        # Details container
        details_frame = ttk.Frame(
            details_window,
            padding=20
        )

        details_frame.pack(
            fill=BOTH,
            expand=True
        )

        details = [
            (
                "Status",
                selected_record.get(
                    "status",
                    "Unknown"
                )
            ),

            (
                "Date & Time",
                selected_record.get(
                    "date_time",
                    ""
                )
            ),

            (
                "Input Folder",
                selected_record.get(
                    "input_folder",
                    ""
                )
            ),

            (
                "Output Folder",
                selected_record.get(
                    "output_folder",
                    ""
                )
            ),

            (
                "Excel Files",
                f"{selected_record.get('excel_files', 0):,}"
            ),

            (
                "Total Rows",
                f"{selected_record.get('total_rows', 0):,}"
            ),

            (
                "Duplicates Removed",
                f"{selected_record.get('duplicates_removed', 0):,}"
            ),

            (
                "Blank Rows Removed",
                f"{selected_record.get('blank_rows_removed', 0):,}"
            ),

            (
                "Final Rows",
                f"{selected_record.get('final_rows', 0):,}"
            ),

            (
                "Processing Time",
                f"{selected_record.get('processing_time', 0):.2f} seconds"
            )
        ]

        for label_text, value_text in details:

            row = ttk.Frame(
                details_frame
            )

            row.pack(
                fill=X,
                pady=6
            )

            ttk.Label(
                row,
                text=f"{label_text}:",
                font=("Segoe UI", 11, "bold"),
                width=22,
                anchor=W
            ).pack(
                side=LEFT
            )

            ttk.Label(
                row,
                text=str(value_text),
                font=("Segoe UI", 11),
                anchor=W
            ).pack(
                side=LEFT,
                fill=X,
                expand=True
            )

        # Close button
        ttk.Button(
            details_window,
            text="✖ Close",
            bootstyle="secondary",
            command=details_window.destroy
        ).pack(
            pady=(10, 25)
        )

    # ==================================================
    # BUTTONS
    # ==================================================

    def create_buttons(self):

        button_frame = ttk.Frame(self.frame)

        button_frame.pack(
            fill=X,
            pady=15
        )

        ttk.Button(
            button_frame,
            text="📂 Open Output Folder",
            bootstyle="success",
            command=self.open_output_folder
        ).pack(
            side=LEFT,
            padx=5
        )

        ttk.Button(
            button_frame,
            text="🗑️ Clear History",
            bootstyle="danger",
            command=self.clear_history
        ).pack(
            side=LEFT,
            padx=5
        )

    # ==================================================
    # LOAD HISTORY
    # ==================================================

    def load_history(self):

        # Clear search
        self.search_var.set("")

        # Clear table
        for item in self.tree.get_children():
            self.tree.delete(item)

        if not HISTORY_FILE.exists():

            self.update_summary([])

            return

        try:

            with open(
                HISTORY_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                history = json.load(file)

        except Exception as e:

            messagebox.showerror(
                "History Error",
                f"Unable to load history:\n\n{e}"
            )

            return

        # Display newest first
        for record in reversed(history):

            self.insert_record(record)

        self.update_summary(history)
#create analytics chart
        self.create_processing_time_chart(history)
        self.create_rows_processed_chart(history)
        return history
        # ==================================================
    # HISTORY ANALYTICS
    # ==================================================

    def calculate_history_analytics(self, history):

        if not history:
            return {
                "total_runs": 0,
                "total_files": 0,
                "total_rows": 0,
                "total_duplicates": 0,
                "average_processing_time": 0,
                "fastest_run": 0,
                "latest_run": ""
            }

        total_runs = len(history)

        total_files = sum(
            int(record.get("excel_files", 0) or 0)
            for record in history
        )

        total_rows = sum(
            int(record.get("total_rows", 0) or 0)
            for record in history
        )

        total_duplicates = sum(
            int(record.get("duplicates_removed", 0) or 0)
            for record in history
        )

        processing_times = []

        for record in history:

            try:
                processing_time = float(
                    record.get("processing_time", 0) or 0
                )

                processing_times.append(processing_time)

            except (TypeError, ValueError):
                continue

        if processing_times:

            average_processing_time = round(
                sum(processing_times) / len(processing_times),
                2
            )

            fastest_run = min(processing_times)

        else:

            average_processing_time = 0
            fastest_run = 0

        latest_run = history[-1].get(
            "date_time",
            ""
        )

        return {
            "total_runs": total_runs,
            "total_files": total_files,
            "total_rows": total_rows,
            "total_duplicates": total_duplicates,
            "average_processing_time": average_processing_time,
            "fastest_run": fastest_run,
            "latest_run": latest_run
        }
        # ==================================================
    # EXPORT HISTORY TO CSV
    # ==================================================

    def export_history_csv(self):

        history = self.load_history()

        if not history:

            messagebox.showinfo(
                "Export History",
                "There is no automation history to export."
            )

            return

        # Ask user where to save the CSV
        file_path = filedialog.asksaveasfilename(
            title="Export Automation History",
            defaultextension=".csv",
            filetypes=[
                (
                    "CSV Files",
                    "*.csv"
                ),
                (
                    "All Files",
                    "*.*"
                )
            ],
            initialfile="automation_history.csv"
        )

        # User cancelled
        if not file_path:
            return

        try:

            fieldnames = [
                "Date & Time",
                "Input Folder",
                "Output Folder",
                "Excel Files",
                "Total Rows",
                "Duplicates Removed",
                "Blank Rows Removed",
                "Final Rows",
                "Processing Time",
                "Status"
            ]

            with open(
                file_path,
                "w",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                writer = csv.DictWriter(
                    file,
                    fieldnames=fieldnames
                )

                writer.writeheader()

                for record in history:

                    writer.writerow({
                        "Date & Time": record.get(
                            "date_time",
                            ""
                        ),
                        "Input Folder": record.get(
                            "input_folder",
                            ""
                        ),
                        "Output Folder": record.get(
                            "output_folder",
                            ""
                        ),
                        "Excel Files": record.get(
                            "excel_files",
                            0
                        ),
                        "Total Rows": record.get(
                            "total_rows",
                            0
                        ),
                        "Duplicates Removed": record.get(
                            "duplicates_removed",
                            0
                        ),
                        "Blank Rows Removed": record.get(
                            "blank_rows_removed",
                            0
                        ),
                        "Final Rows": record.get(
                            "final_rows",
                            0
                        ),
                        "Processing Time": record.get(
                            "processing_time",
                            0
                        ),
                        "Status": record.get(
                            "status",
                            ""
                        )
                    })

            messagebox.showinfo(
                "Export Successful",
                f"Automation history exported successfully.\n\n"
                f"File:\n{file_path}"
            )

        except Exception as e:

            messagebox.showerror(
                "Export Failed",
                f"Unable to export automation history.\n\n"
                f"Error: {e}"
            )
        # ==================================================
    # EXPORT HISTORY TO EXCEL
    # ==================================================

    def export_history_excel(self):

        history = self.load_history()

        if not history:

            messagebox.showinfo(
                "Export History",
                "There is no automation history to export."
            )

            return

        file_path = filedialog.asksaveasfilename(
            title="Export Automation History to Excel",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Files", "*.xlsx"),
                ("All Files", "*.*")
            ],
            initialfile="automation_history.xlsx"
        )

        if not file_path:
            return

        try:

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Automation History"

            headers = [
                "Date & Time",
                "Input Folder",
                "Output Folder",
                "Excel Files",
                "Total Rows",
                "Duplicates Removed",
                "Blank Rows Removed",
                "Final Rows",
                "Processing Time",
                "Status"
            ]

            worksheet.append(headers)

            # Header formatting
            for cell in worksheet[1]:

                cell.font = Font(bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            # Add history records
            for record in history:

                worksheet.append([
                    record.get("date_time", ""),
                    record.get("input_folder", ""),
                    record.get("output_folder", ""),
                    record.get("excel_files", 0),
                    record.get("total_rows", 0),
                    record.get("duplicates_removed", 0),
                    record.get("blank_rows_removed", 0),
                    record.get("final_rows", 0),
                    record.get("processing_time", 0),
                    record.get("status", "")
                ])

            # Alignment
            for row in worksheet.iter_rows():

                for cell in row:

                    cell.alignment = Alignment(
                        vertical="center"
                    )

            # Center numeric/status columns
            for row in worksheet.iter_rows(
                min_row=2,
                min_col=4,
                max_col=10
            ):

                for cell in row:

                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center"
                    )

            # Auto-size columns
            for column_cells in worksheet.columns:

                max_length = 0
                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    60
                )

            # Freeze header row
            worksheet.freeze_panes = "A2"

            # Add Excel filter
            worksheet.auto_filter.ref = worksheet.dimensions

            workbook.save(file_path)

            messagebox.showinfo(
                "Export Successful",
                f"Automation history exported successfully.\n\n"
                f"File:\n{file_path}"
            )

        except Exception as e:

            messagebox.showerror(
                "Export Failed",
                f"Unable to export automation history.\n\n"
                f"Error: {e}"
            )
   

# ==================================================
# UPDATE SUMMARY
# ==================================================

    def update_summary(self, history):

        total_runs = len(history)

        successful_runs = sum(
            1
            for record in history
            if record.get("status") == "Success"
        )

        total_files = sum(
            int(record.get("excel_files", 0) or 0)
            for record in history
        )

        total_rows = sum(
            int(record.get("total_rows", 0) or 0)
            for record in history
        )

        # ==================================================
        # HISTORY ANALYTICS
        # ==================================================

        analytics = self.calculate_history_analytics(
            history
        )

        # ==================================================
        # EXISTING SUMMARY CARDS
        # ==================================================

        self.total_runs_label.config(
            text=f"{total_runs:,}"
        )

        self.success_label.config(
            text=f"{successful_runs:,}"
        )

        self.files_label.config(
            text=f"{total_files:,}"
        )

        self.rows_label.config(
            text=f"{total_rows:,}"
        )

        # ==================================================
        # ANALYTICS CARDS
        # ==================================================

        self.avg_time_label.config(
            text=f"{analytics['average_processing_time']:.2f} sec"
        )

        self.fastest_run_label.config(
            text=f"{analytics['fastest_run']:.2f} sec"
        )

        self.duplicates_label.config(
            text=f"{analytics['total_duplicates']:,}"
        )

    # ==================================================
    # OPEN OUTPUT FOLDER
    # ==================================================

    def open_output_folder(self):

        output_folder = self.app.settings.get(
            "output_folder",
            ""
        )

        if not output_folder:

            messagebox.showwarning(
                "Output Folder",
                "No output folder has been configured."
            )

            return

        folder = Path(output_folder)

        if not folder.exists():

            messagebox.showwarning(
                "Output Folder",
                "The output folder does not exist."
            )

            return

        try:

            os.startfile(folder)

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"Unable to open folder:\n\n{e}"
            )

    # ==================================================
    # CLEAR HISTORY
    # ==================================================

    def clear_history(self):

        if not HISTORY_FILE.exists():

            messagebox.showinfo(
                "History",
                "There is no history to clear."
            )

            return

        confirm = messagebox.askyesno(
            "Clear History",
            "Are you sure you want to delete all automation history?"
        )

        if not confirm:
            return

        try:

            with open(
                HISTORY_FILE,
                "w",
                encoding="utf-8"
            ) as file:

                json.dump(
                    [],
                    file,
                    indent=4
                )

            self.load_history()

            messagebox.showinfo(
                "History",
                "Automation history cleared successfully."
            )

        except Exception as e:

            messagebox.showerror(
                "History Error",
                f"Unable to clear history:\n\n{e}"
            )
# ==================================================
# PROCESSING TIME CHART
# ==================================================

    def create_processing_time_chart(self, history):
    #remove previous chart
        if hasattr(self,"processing_chart_frame"):
            try:
                self.processing_chart_figure.clear()
                self.processing_chart_frame.destroy()
            except Exception:
                pass
        if not history:
            return

        chart_frame = ttk.Frame(self.frame)

        chart_frame.pack(
            fill=X,
            pady=(0, 20)
        )

        # Keep only the latest 10 runs
        chart_history = history[-10:]

        labels = []
        processing_times = []

        for index, record in enumerate(chart_history, start=1):

            labels.append(
                f"Run {index}"
            )

            try:
                processing_time = float(
                    record.get("processing_time", 0) or 0
                )
            except (TypeError, ValueError):
                processing_time = 0

            processing_times.append(
                processing_time
            )

        figure, axis = plt.subplots(
            figsize=(10, 3.2)
        )

        axis.plot(
            labels,
            processing_times,
            marker="o"
        )

        axis.set_title(
            "Processing Time per Run"
        )

        axis.set_xlabel(
            "Automation Run"
        )

        axis.set_ylabel(
            "Processing Time (seconds)"
        )

        axis.grid(
            True,
            alpha=0.3
        )

        figure.tight_layout()

        canvas = FigureCanvasTkAgg(
            figure,
            master=chart_frame
        )

        canvas.draw()

        canvas.get_tk_widget().pack(
            fill=BOTH,
            expand=True
        )

        # Store references so the chart can be managed later
        self.processing_chart_frame = chart_frame
        self.processing_chart_figure = figure
        self.processing_chart_canvas = canvas
        # ==================================================
    # ROWS PROCESSED CHART
    # ==================================================

    def create_rows_processed_chart(self, history):

        # Remove previous chart
        if hasattr(self, "rows_chart_frame"):

            try:
                self.rows_chart_figure.clear()
                self.rows_chart_frame.destroy()
            except Exception:
                pass

        if not history:
            return

        chart_frame = ttk.Frame(self.frame)

        chart_frame.pack(
            fill=X,
            pady=(0, 20)
        )

        # Keep only the latest 10 runs
        chart_history = history[-10:]

        labels = []
        rows_processed = []

        for index, record in enumerate(chart_history, start=1):

            labels.append(
                f"Run {index}"
            )

            try:
                rows = int(
                    record.get("total_rows", 0) or 0
                )
            except (TypeError, ValueError):
                rows = 0

            rows_processed.append(
                rows
            )

        figure, axis = plt.subplots(
            figsize=(10, 3.2)
        )

        axis.plot(
            labels,
            rows_processed,
            marker="o"
        )

        axis.set_title(
            "Rows Processed per Run"
        )

        axis.set_xlabel(
            "Automation Run"
        )

        axis.set_ylabel(
            "Rows Processed"
        )

        axis.grid(
            True,
            alpha=0.3
        )

        figure.tight_layout()

        canvas = FigureCanvasTkAgg(
            figure,
            master=chart_frame
        )

        canvas.draw()

        canvas.get_tk_widget().pack(
            fill=BOTH,
            expand=True
        )

        # Store references
        self.rows_chart_frame = chart_frame
        self.rows_chart_figure = figure
        self.rows_chart_canvas = canvas