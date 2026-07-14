from openpyxl import load_workbook

def create_summary(file_path,stats):
    wb=load_workbook(file_path)
    
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
        
    ws=wb.create_sheet("Summary")
    
    ws["A1"]="Excel Automation Toolkit"
    ws["A3"]="REport"
    ws["B3"]="Value"
    
    row=4
    for key,value in stats.items():
        ws.cell(row=row,column=1).value=key
        ws.cell(row=row,column=2).value=value
        row+=1
        
    wb.save(file_path)