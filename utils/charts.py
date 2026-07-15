from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def create_chart(file_path):
    wb=load_workbook(file_path)
    ws=wb["Summary"]
    chart=BarChart()
    data=Reference(ws,
                   min_col=2,
                   min_row=4,
                   max_row=7)
    categories=Reference(ws,
                         min_col=1,
                         min_row=4,
                         max_row=7)
    
    chart.add_data(data)
    chart.set_categories(categories)
    chart.title="Cleaning Report"
    chart.y_axis.title="Count"
    chart.x_axis.title="Metrics"
    ws.add_chart(chart,"E3")
    
    wb.save(file_path)
    