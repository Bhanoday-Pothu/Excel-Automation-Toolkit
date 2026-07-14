from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Side

def format_excel(file_path):
    wb=load_workbook(file_path)
    ws=wb.active
    
    header_fill=PatternFill(
        fill_type="solid",
        start_color="4F81BD",
        end_color="4F81BD"
    )
    header_font= Font(
        bold=True,
        color="FFFFFF"
    )
    
    thin=Side(style="thin")
    
    border=Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )
    
    for cell in ws[1]:
        cell.fill=header_fill
        cell.font=header_font
        cell.alignment=Alignment(horizontal="center")
        
    
    for row in ws.iter_rows():
        for cell in row:
            cell.border=border
            
    for column in ws.columns:
        max_length=0
        column_letter=column[0].column_letter
        
        for cell in column:
            try:
                if len(str(cell.value))>max_length:
                    max_length = len(str(cell.value))
                    
            except:
                pass
            
        adjusted_width=max_length+5
        ws.column_dimensions[column_letter].width=adjusted_width
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    wb.save(file_path)