from openpyxl import Workbook, load_workbook

from utils.summary import create_summary


def test_create_summary(tmp_path):
    file_path = tmp_path / "sales.xlsx"

    wb = Workbook()
    ws = wb.active

    ws["A1"] = "Name"
    ws["B1"] = "Amount"
    ws["A2"] = "Alice"
    ws["B2"] = 50000

    wb.save(file_path)
    wb.close()

    stats = {
        "Total Rows": 10,
        "Duplicate Rows Removed": 2,
        "Blank Rows Removed": 1,
        "Final Rows": 7,
    }

    result = create_summary(file_path, stats)

    assert result is None

    formatted_wb = load_workbook(file_path)
    assert "Summary" in formatted_wb.sheetnames

    summary_ws = formatted_wb["Summary"]

    assert summary_ws["A1"].value == "Excel Automation Toolkit"
    assert summary_ws["A3"].value == "REport"
    assert summary_ws["B3"].value == "Value"

    assert summary_ws["A4"].value == "Total Rows"
    assert summary_ws["B4"].value == 10

    assert summary_ws["A5"].value == "Duplicate Rows Removed"
    assert summary_ws["B5"].value == 2

    assert summary_ws["A6"].value == "Blank Rows Removed"
    assert summary_ws["B6"].value == 1

    assert summary_ws["A7"].value == "Final Rows"
    assert summary_ws["B7"].value == 7

    formatted_wb.close()


def test_create_summary_replaces_existing_summary(tmp_path):
    file_path = tmp_path / "sales.xlsx"

    wb = Workbook()
    ws = wb.active

    ws["A1"] = "Original Data"

    old_summary = wb.create_sheet("Summary")
    old_summary["A1"] = "Old Summary"
    old_summary["A10"] = "Old Data"

    wb.save(file_path)
    wb.close()

    stats = {
        "Total Rows": 5,
        "Final Rows": 5,
    }

    create_summary(file_path, stats)

    formatted_wb = load_workbook(file_path)

    assert "Summary" in formatted_wb.sheetnames

    summary_ws = formatted_wb["Summary"]

    assert summary_ws["A1"].value == "Excel Automation Toolkit"
    assert summary_ws["A10"].value is None
    assert summary_ws["A4"].value == "Total Rows"
    assert summary_ws["B4"].value == 5

    formatted_wb.close()