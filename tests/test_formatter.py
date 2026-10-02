from openpyxl import Workbook, load_workbook

from utils.formatter import format_excel


def test_format_excel(tmp_path):
    file_path = tmp_path / "sales.xlsx"

    wb = Workbook()
    ws = wb.active

    ws.append(["Name", "Product", "Amount", "City"])
    ws.append(["Alice", "Laptop", 50000, "Hyderabad"])
    ws.append(["Bob", "Mouse", 1000, "Bangalore"])

    wb.save(file_path)

    result = format_excel(file_path)

    assert result is None

    formatted_wb = load_workbook(file_path)
    formatted_ws = formatted_wb.active

    # Header formatting
    for cell in formatted_ws[1]:
        assert cell.font.bold is True
        assert cell.font.color.rgb == "00FFFFFF"
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb == "004F81BD"
        assert cell.alignment.horizontal == "center"

    # Borders
    for row in formatted_ws.iter_rows():
        for cell in row:
            assert cell.border.left.style == "thin"
            assert cell.border.right.style == "thin"
            assert cell.border.top.style == "thin"
            assert cell.border.bottom.style == "thin"

    # Freeze panes
    assert formatted_ws.freeze_panes == "A2"

    # AutoFilter
    assert formatted_ws.auto_filter.ref == formatted_ws.dimensions

    # Column widths should have been adjusted
    assert formatted_ws.column_dimensions["A"].width > 0
    assert formatted_ws.column_dimensions["B"].width > 0
    assert formatted_ws.column_dimensions["C"].width > 0
    assert formatted_ws.column_dimensions["D"].width > 0

    formatted_wb.close()