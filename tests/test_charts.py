from openpyxl import Workbook, load_workbook

from utils.charts import create_chart


def test_create_chart(tmp_path):
    file_path = tmp_path / "sales.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Excel Automation Toolkit"
    ws["A3"] = "REport"
    ws["B3"] = "Value"

    stats = [
        ("Total Rows", 10),
        ("Duplicate Rows Removed", 2),
        ("Blank Rows Removed", 1),
        ("Final Rows", 7),
    ]

    for row_number, (metric, value) in enumerate(stats, start=4):
        ws.cell(row=row_number, column=1).value = metric
        ws.cell(row=row_number, column=2).value = value

    wb.save(file_path)
    wb.close()

    result = create_chart(file_path)

    assert result is None

    formatted_wb = load_workbook(file_path)
    summary_ws = formatted_wb["Summary"]

    # One chart should have been added.
    assert len(summary_ws._charts) == 1

    chart = summary_ws._charts[0]

    # Verify chart configuration.
    assert chart.title.tx.rich.p[0].r[0].t == "Cleaning Report"
    assert chart.y_axis.title.tx.rich.p[0].r[0].t == "Count"
    assert chart.x_axis.title.tx.rich.p[0].r[0].t == "Metrics"

    # Verify chart position.
    assert chart.anchor._from.col == 4
    assert chart.anchor._from.row == 2

    formatted_wb.close()